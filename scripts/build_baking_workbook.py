"""Baking-specific, explicitly heuristic action layer; original evidence stays intact."""
import argparse
import os
import re
from pathlib import Path
from tempfile import NamedTemporaryFile
from openpyxl import load_workbook
from build_workbook import build, load_capture, add_sheet


def classify(term):
    rules = [
        (r'招聘|面试|工资|工厂改造', '非零售需求', 1, 1, '不优先用于商品运营'),
        (r'代理|加盟|经销', '渠道合作', 8, 3, '人工确认渠道合作条件，不当作零售买家'),
        (r'上涨原因|涨价原因|行情', '行情信息', 4, 7, '核实行情与日期后再写，不把询价等同采购'),
        (r'在哪里买|哪里买|购买|下单|旗舰店|正品店', '购买渠道', 10, 6, '补充真实购买渠道、规格及发货说明'),
        (r'批发|供货|供应商|销售电话', '批量采购', 9, 6, '补充起订量、包装、运费及报价条件'),
        (r'不出膜|失败|塌|开裂|不起|发不|揉不|太硬', '制作问题', 4, 9, '整理排查步骤，区分原料、配方和操作原因'),
        (r'区别|对比|测评|哪个牌子|哪种好|哪个好|怎么选', '选品比较', 7, 9, '用可核实的用途、参数、规格做比较'),
        (r'多少钱|价格|到手价|优惠', '价格规格', 8, 6, '明确规格及价格条件，不编造优惠'),
        (r'蛋白|吸水|灰分|筋度|配料|保质|储存', '原料参数', 6, 8, '依据包装或厂家资料补充参数与保存说明'),
        (r'配方|教程|做法|比例|怎么做', '配方教程', 4, 9, '提供经过验证的配方与操作条件'),
        (r'吐司|面包|贝果|披萨|包子|馒头|蛋挞|饼干|蛋糕|面条', '用途适配', 6, 9, '说明适用用途及限制，避免承诺所有配方通用'),
        (r'25\s*(kg|公斤|千克)|商用|整袋', '商用规格线索', 7, 5, '核实包装与供货条件，不能凭词确认买家身份'),
    ]
    category, purchase, content, action = '待人工判断', 4, 4, '结合商品与搜索场景人工判断'
    for pattern, c, p, v, a in rules:
        if re.search(pattern, term, re.I):
            category, purchase, content, action = c, p, v, a
            break
    audience = '未明确'
    if re.search(r'批发|商用|25\s*(kg|公斤|千克)|整袋', term, re.I):
        audience = '批量/商用线索（非身份结论）'
    elif re.search(r'家用|家庭|新手', term):
        audience = '家庭/新手线索（非身份结论）'
    uses = '、'.join(x for x in ['吐司','面包','贝果','披萨','包子','馒头','蛋挞','饼干','蛋糕','面条'] if x in term)
    return category, purchase, content, audience, uses, action


def enhance(source, output, classifier=classify, category_sheet='烘焙需求分类', version='baking-v1-20260930'):
    source, output = Path(source), Path(output)
    if output.exists():
        raise ValueError('请使用新的输出文件名，避免覆盖原表')
    data, roots, queries = load_capture(source)
    for q in queries.values():
        ranks = sorted(s['rank'] for s in q['suggestions'])
        if ranks != list(range(1, len(ranks) + 1)):
            raise ValueError('下拉位置不连续，请核对截图: ' + q['query'])
    counts = build(source, output)
    wb = load_workbook(output)
    classified, actions, cross = [], [], []
    for term, original, reason, paths, root_text in wb['去重词库'].iter_rows(min_row=2, values_only=True):
        category, purchase, content, audience, uses, action = classifier(term)
        classified.append([term, original, purchase, content, category, audience, uses, action, root_text, '规则推断；不是搜索量或成交预测'])
        cross.append([term, root_text, paths, '出现路径数不是热度；多个根词共用建议也不是独立用户数'])
        if category not in ('非零售需求', '待人工判断', '渠道合作'):
            for kind in ('详情页检查', '内容选题', '客服答疑'):
                suggestion = {'详情页检查': '检查是否已回答：' + term,
                              '内容选题': '围绕“' + term + '”制作实测或解释内容',
                              '客服答疑': '整理“' + term + '”的准确回答'}[kind]
                actions.append([term, kind, suggestion, action, '确认本店商品适配；核实参数/价格/配方后使用', '规则生成建议，不是采集原文'])
    add_sheet(wb, category_sheet, ['真实建议词','原版商业分','采购意图分','内容价值分','需求分类','受众线索','用途线索','评分与行动依据','来源根词','边界'], classified)
    add_sheet(wb, '商家行动建议', ['真实建议词','用途','建议动作','依据','发布前核实','来源类型'], actions)
    add_sheet(wb, '品牌品类对照', ['真实建议词','来源根词','路径数','注意'], cross)
    add_sheet(wb, '增强说明', ['项目','说明'], [
        ['原始证据', '原五张表、原评分和截图路径保留，不修改采集 JSON'],
        ['采购意图', '1–10 分，依据词面规则估计接近采购的程度；非转化率'],
        ['内容价值', '1–10 分，依据用途、配方、选品与排错需求估计解答价值；非流量预测'],
        ['平台联想词', '只能证明采集时平台展示过；不证明高频、热度、销量或竞争程度'],
        ['使用方式', '人工复核分类；结合商品资料，再修改详情页、制作内容或整理客服答案'],
        ['规则版本', version],
    ])
    temp = None
    try:
        with NamedTemporaryFile(dir=output.parent, suffix='.xlsx', delete=False) as handle:
            temp = Path(handle.name)
        wb.save(temp)
        os.replace(temp, output)
    finally:
        if temp:
            temp.unlink(missing_ok=True)
    return counts


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    print('路径、去重词、待采集:', enhance(args.input, args.output))
